# Converte a base de dados de níveis (Excel do caderno pedagógico) em scripts/niveis/niveis-cfp.json.
# Uso: python3 scripts/niveis/xlsx-to-json.py <base_dados_niveis_CFP.xlsx>   (requer openpyxl)
import json
import sys
from pathlib import Path

import openpyxl

OUT = Path(__file__).with_name("niveis-cfp.json")


def clean(value):
    if value is None:
        return None
    text = str(value).strip()
    return None if text in ("", "—") else text


def split_list(value, sep):
    text = clean(value)
    return [part.strip() for part in text.split(sep)] if text else []


def to_int(value):
    digits = "".join(c for c in str(value or "") if c.isdigit())
    return int(digits) if digits else None


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)

    levels = []
    for i, row in enumerate(wb["Visão Geral"].iter_rows(min_row=2, values_only=True), start=1):
        if not clean(row[0]):
            continue
        code, stage, age, duration, space, capacity, requirements = row[:7]
        levels.append({
            "code": clean(code),
            "stage": clean(stage),
            "age_range": clean(age),
            "duration_min": to_int(duration),
            "space": clean(space),
            "max_students": to_int(capacity),
            "requirements": clean(requirements),
            "periods": split_list(row[8], "·"),
            "order_index": i,
        })

    contents = []
    order = {}
    for row in wb["Base de Dados"].iter_rows(min_row=2, values_only=True):
        if not clean(row[0]):
            continue
        level = clean(row[0])
        order[level] = order.get(level, 0) + 1
        contents.append({
            "level_code": level,
            "number": clean(row[6]),
            "block": clean(row[7]),
            "title": clean(row[8]),
            "intro_period": clean(row[9]),
            "periods": split_list(row[10], "+"),
            "start_month": clean(row[11]),
            "months": split_list(row[12], ","),
            "continuity": split_list(row[13], "+"),
            "source": clean(row[14]),
            "order_index": order[level],
        })

    OUT.write_text(json.dumps({"levels": levels, "contents": contents}, ensure_ascii=False, indent=2) + "\n")
    print(f"{len(levels)} níveis, {len(contents)} conteúdos -> {OUT}")


if __name__ == "__main__":
    main(sys.argv[1])
